"""Conversions between PDF and other formats, plus OCR.

Every function takes ``pathlib.Path`` arguments and reports progress through an
optional ``progress(percent, message)`` callback. The callback may raise
``OperationCancelled`` to stop a conversion; partial output is removed before the
exception propagates. Errors that users should see are raised as
``PdfStudioError`` (or ``MissingDependencyError`` / ``PasswordRequiredError``).

External programs are looked up on PATH and in the usual Windows install folders.
Windows-only Microsoft Office support uses pywin32 and is imported lazily.
"""

from __future__ import annotations

import codecs
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from collections.abc import Callable, Iterator, Sequence
from contextlib import contextmanager, suppress
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any

import pymupdf
from PIL import Image, ImageOps

from pdfstudio.core.errors import (
    MissingDependencyError,
    PasswordRequiredError,
    PdfStudioError,
)
from pdfstudio.core.files import unique_path

ProgressCallback = Callable[[int, str], None]

_LOG = logging.getLogger(__name__)

_DPI_MIN = 36
_DPI_MAX = 600
_JPEG_QUALITY = 90
_PT_PER_PX = 72.0 / 96.0  # Browsers and Office treat 96 px as one inch
_MAX_PAGE_PT = 14400.0  # Largest page dimension the PDF format allows (200 inches)
_EMU_PER_PT = 12700
_IMAGE_SUFFIXES = frozenset({".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff", ".webp"})
_ALPHA_MODES = frozenset({"RGBA", "RGBa", "LA", "La", "PA"})
_ILLEGAL_TEXT_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")
_LANGUAGE_CODE = re.compile(r"[A-Za-z0-9_+\-]+")
_LATIN_WORD = re.compile(r"[A-Za-z0-9]{2,}")
_ARABIC_CHARS = re.compile("[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff\ufb50-\ufdff\ufe70-\ufeff]")
_META_CHARSET = re.compile(rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9_.:-]+)""", re.IGNORECASE)
_BASE_FONT_FILE = "pdfstudio-base.ttf"
_MAX_HTML_PAGES = 5000
_LIBREOFFICE_TIMEOUT = 180
_TESSERACT_PAGE_TIMEOUT = 600

_OFFICE_APPS = {
    ".doc": "word",
    ".docx": "word",
    ".odt": "word",
    ".rtf": "word",
    ".xls": "excel",
    ".xlsx": "excel",
    ".ods": "excel",
    ".ppt": "powerpoint",
    ".pptx": "powerpoint",
    ".odp": "powerpoint",
}
_OFFICE_PROG_IDS = {
    "word": "Word.Application",
    "excel": "Excel.Application",
    "powerpoint": "PowerPoint.Application",
}
_OFFICE_LABELS = {
    "word": "Microsoft Word",
    "excel": "Microsoft Excel",
    "powerpoint": "Microsoft PowerPoint",
}

_OFFICE_MISSING = (
    "This conversion needs Microsoft Office (Word, Excel or PowerPoint) or LibreOffice "
    "installed on this computer."
)
_TESSERACT_MISSING = (
    "OCR needs Tesseract OCR. Install it from https://github.com/UB-Mannheim/tesseract/wiki "
    "and try again."
)


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------


def _report(progress: ProgressCallback | None, percent: float, message: str) -> None:
    if progress is not None:
        progress(max(0, min(100, int(percent))), message)


def _share(done: int, total: int, start: float = 0.0, span: float = 100.0) -> float:
    """Percentage at ``done`` of ``total`` steps, mapped into ``[start, start + span]``."""
    if total <= 0:
        return start + span
    return start + span * done / total


def _check_dpi(dpi: Any) -> int:
    try:
        value = int(dpi)
    except (TypeError, ValueError) as exc:
        raise PdfStudioError("The resolution must be a number of dpi.") from exc
    if not _DPI_MIN <= value <= _DPI_MAX:
        raise PdfStudioError(f"The resolution must be between {_DPI_MIN} and {_DPI_MAX} dpi.")
    return value


def _no_window_kwargs() -> dict[str, int]:
    """Keep console programs from flashing a window on Windows."""
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    return {"creationflags": flags} if flags else {}


def _discard(paths: Sequence[Path]) -> None:
    for path in paths:
        with suppress(OSError):
            Path(path).unlink()


@contextmanager
def _remove_partial_output(target: Path) -> Iterator[None]:
    """Delete ``target`` if the wrapped block fails and this call created it."""
    existed = target.exists()
    try:
        yield
    except BaseException:
        if not existed:
            _discard([target])
        raise


def _open_pdf(path: Path, password: str | None) -> pymupdf.Document:
    """Open a PDF, authenticating with ``password`` when it is encrypted."""
    if not path.is_file():
        raise PdfStudioError(f"'{path.name}' was not found.")
    try:
        document = pymupdf.open(str(path))
    except Exception as exc:
        raise PdfStudioError(
            f"'{path.name}' could not be opened. The file may be damaged or not a PDF."
        ) from exc
    try:
        if not document.is_pdf:
            raise PdfStudioError(f"'{path.name}' is not a PDF file.")
        if document.needs_pass and not document.authenticate(password or ""):
            if password:
                raise PasswordRequiredError(f"The password for '{path.name}' is not correct.")
            raise PasswordRequiredError(
                f"'{path.name}' is password protected. Enter its password and try again."
            )
        if document.page_count == 0:
            raise PdfStudioError(f"'{path.name}' has no pages.")
    except BaseException:
        document.close()
        raise
    return document


def _resolve_pages(pages: Sequence[int] | None, page_count: int) -> list[int]:
    """Validate 0-based page indexes; ``None`` means every page."""
    if pages is None:
        return list(range(page_count))
    indexes: list[int] = []
    for raw in pages:
        index = int(raw)
        if not 0 <= index < page_count:
            plural = "s" if page_count != 1 else ""
            raise PdfStudioError(
                f"Page {index + 1} does not exist. This PDF has {page_count} page{plural}."
            )
        if index not in indexes:
            indexes.append(index)
    if not indexes:
        raise PdfStudioError("Choose at least one page.")
    return indexes


def _render_pixmap(page: pymupdf.Page, dpi: int) -> pymupdf.Pixmap:
    # alpha=False composites the page onto a white background, so transparent
    # areas of the PDF come out white instead of black or see-through.
    return page.get_pixmap(dpi=dpi, alpha=False)


def _render_jpeg(page: pymupdf.Page, dpi: int) -> bytes:
    return _render_pixmap(page, dpi).tobytes("jpg", jpg_quality=_JPEG_QUALITY)


# ---------------------------------------------------------------------------
# PDF to images
# ---------------------------------------------------------------------------


def pdf_to_images(
    path: Path,
    out_dir: Path,
    stem: str,
    fmt: str = "jpg",
    dpi: int = 150,
    pages: Sequence[int] | None = None,
    password: str | None = None,
    progress: ProgressCallback | None = None,
) -> list[Path]:
    """Render pages to one JPG or PNG file each.

    Files are named ``"<stem> - page 001.jpg"`` using the 1-based page number, so a
    subset of pages keeps the numbers of the original document. Existing files are
    never overwritten; ``" (2)"`` is added to the name instead.
    """
    source = Path(path)
    fmt_key = str(fmt).lower()
    if fmt_key not in ("jpg", "png"):
        raise PdfStudioError("The image format must be JPG or PNG.")
    dpi = _check_dpi(dpi)
    document = _open_pdf(source, password)
    created: list[Path] = []
    try:
        indexes = _resolve_pages(pages, document.page_count)
        total = len(indexes)
        for step, index in enumerate(indexes):
            _report(progress, _share(step, total), f"Rendering page {index + 1} of {total}")
            pixmap = _render_pixmap(document[index], dpi)
            if fmt_key == "jpg":
                data = pixmap.tobytes("jpg", jpg_quality=_JPEG_QUALITY)
            else:
                data = pixmap.tobytes("png")
            target = unique_path(Path(out_dir), f"{stem} - page {index + 1:03d}", fmt_key)
            target.write_bytes(data)
            created.append(target)
        _report(progress, 100, "Done")
    except BaseException:
        _discard(created)
        raise
    finally:
        document.close()
    return created


# ---------------------------------------------------------------------------
# Images to PDF
# ---------------------------------------------------------------------------


def images_to_pdf(
    images: Sequence[Path],
    output: Path,
    page_size: str = "fit",
    margin: float = 0,
    progress: ProgressCallback | None = None,
) -> Path:
    """Combine images into one PDF with one page per image (per frame for multi-page TIFF).

    ``page_size="fit"`` makes each page exactly the image size (96 px = 72 pt), so
    nothing is resampled. ``"a4"`` and ``"letter"`` scale the image to fit inside the
    page minus ``margin`` points, centred, keeping its aspect ratio; a landscape image
    gets a landscape page. JPEG and PNG files are embedded unchanged when no conversion
    is needed (no alpha, CMYK, 16-bit or EXIF rotation). Anything else is converted to
    RGB or grey first: JPEG in, JPEG out at quality 95; other formats become PNG.
    """
    paths = [Path(item) for item in images]
    if not paths:
        raise PdfStudioError("Choose at least one image to convert.")
    size_key = str(page_size).lower()
    if size_key not in ("fit", "a4", "letter"):
        raise PdfStudioError("The page size must be Fit to image, A4 or Letter.")
    margin = float(margin)
    if margin < 0:
        raise PdfStudioError("The margin cannot be negative.")
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)

    document = pymupdf.open()
    try:
        total = len(paths)
        for number, image_path in enumerate(paths, start=1):
            _report(progress, _share(number - 1, total, 0, 90), f"Adding image {number} of {total}: {image_path.name}")
            for data, width_px, height_px in _image_frames(image_path):
                _add_image_page(document, data, width_px, height_px, size_key, margin)
        _report(progress, 92, "Saving PDF")
        with _remove_partial_output(target):
            document.save(str(target), garbage=3, deflate=True)
    finally:
        document.close()
    _report(progress, 100, "Done")
    return target


def _image_frames(path: Path) -> list[tuple[bytes, int, int]]:
    """Return ``(embeddable_bytes, width_px, height_px)`` for every frame of an image."""
    if path.suffix.lower() not in _IMAGE_SUFFIXES:
        raise PdfStudioError(
            f"'{path.name}' is not a supported image type. Use JPG, PNG, BMP, GIF, TIFF or WebP."
        )
    if not path.is_file():
        raise PdfStudioError(f"'{path.name}' was not found.")
    try:
        with Image.open(path) as image:
            frame_count = getattr(image, "n_frames", 1)
            frames: list[tuple[bytes, int, int]] = []
            for index in range(frame_count):
                if index:
                    image.seek(index)
                frames.append(_embeddable(image, path, single=frame_count == 1))
            return frames
    except Exception as exc:
        raise PdfStudioError(
            f"'{path.name}' could not be read as an image. The file may be damaged."
        ) from exc


def _embeddable(image: Image.Image, path: Path, single: bool) -> tuple[bytes, int, int]:
    fmt = image.format
    orientation = _exif_orientation(image)
    if (
        single
        and orientation == 1
        and fmt in ("JPEG", "PNG")
        and image.mode in ("RGB", "L")
        and "transparency" not in image.info
    ):
        return path.read_bytes(), image.width, image.height

    frame = ImageOps.exif_transpose(image) if orientation != 1 else image
    flat = _flatten(frame)
    buffer = BytesIO()
    if fmt == "JPEG":
        flat.save(buffer, format="JPEG", quality=95)
    else:
        flat.save(buffer, format="PNG")
    return buffer.getvalue(), flat.width, flat.height


def _exif_orientation(image: Image.Image) -> int:
    try:
        return int(image.getexif().get(0x0112, 1) or 1)
    except Exception:
        return 1


def _flatten(image: Image.Image) -> Image.Image:
    """Return an 8-bit RGB or grey image with transparency composited onto white."""
    mode = image.mode
    if mode.startswith("I") or mode == "F":
        return _to_8bit_grey(image)
    if mode in _ALPHA_MODES or "transparency" in image.info:
        rgba = image.convert("RGBA")
        background = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        return Image.alpha_composite(background, rgba).convert("RGB")
    if mode in ("RGB", "L"):
        return image
    if mode == "1":
        return image.convert("L")
    return image.convert("RGB")  # P, CMYK, YCbCr, LAB, HSV


def _to_8bit_grey(image: Image.Image) -> Image.Image:
    values = image if image.mode in ("I", "F") else image.convert("I")
    _, high = values.getextrema()
    if values.mode == "I" and high <= 255:
        return values.convert("L")
    scale = 255.0 / high if high > 0 else 1.0
    return values.point(lambda value: value * scale).convert("L")


def _add_image_page(
    document: pymupdf.Document,
    data: bytes,
    width_px: int,
    height_px: int,
    size_key: str,
    margin: float,
) -> None:
    width_pt = width_px * _PT_PER_PX
    height_pt = height_px * _PT_PER_PX
    if size_key == "fit":
        scale = min(1.0, _MAX_PAGE_PT / max(width_pt, height_pt))
        page_w, page_h = width_pt * scale, height_pt * scale
        image_w, image_h = page_w, page_h
    else:
        paper = pymupdf.paper_rect(size_key)
        page_w, page_h = paper.width, paper.height
        if width_px > height_px:
            page_w, page_h = page_h, page_w
        avail_w = page_w - 2 * margin
        avail_h = page_h - 2 * margin
        if avail_w <= 0 or avail_h <= 0:
            raise PdfStudioError("The margin is too large for the page size.")
        scale = min(avail_w / width_pt, avail_h / height_pt)
        image_w, image_h = width_pt * scale, height_pt * scale
    page = document.new_page(width=page_w, height=page_h)
    x0 = (page_w - image_w) / 2
    y0 = (page_h - image_h) / 2
    rect = pymupdf.Rect(x0, y0, x0 + image_w, y0 + image_h)
    page.insert_image(rect, stream=data, keep_proportion=True)


# ---------------------------------------------------------------------------
# PDF to Word
# ---------------------------------------------------------------------------


def pdf_to_docx(
    path: Path,
    output: Path,
    pages: Sequence[int] | None = None,
    password: str | None = None,
    progress: ProgressCallback | None = None,
) -> Path:
    """Convert to .docx with pdf2docx (layout, text, tables and images).

    The steps of ``Converter.convert`` are driven one by one so progress can be
    reported per page. Like pdf2docx's default, a page it cannot parse is skipped
    and logged instead of failing the whole conversion.
    """
    from pdf2docx import Converter

    source = Path(path)
    target = Path(output)
    document = _open_pdf(source, password)
    try:
        indexes = _resolve_pages(pages, document.page_count)
    finally:
        document.close()

    _report(progress, 2, "Analyzing the PDF")
    converter: Converter | None = None
    with _remove_partial_output(target):
        try:
            converter = Converter(str(source), password=password or None)
            settings = converter.default_settings
            converter.load_pages(pages=None if pages is None else indexes)
            converter.parse_document(**settings)

            to_parse = [page for page in converter.pages if not page.skip_parsing]
            total = len(to_parse)
            for number, page in enumerate(to_parse, start=1):
                _report(progress, _share(number - 1, total, 5, 75), f"Converting page {page.id + 1} ({number} of {total})")
                try:
                    page.parse(**settings)
                except Exception as exc:
                    if settings.get("raw_exceptions") or settings.get("debug") or not settings.get("ignore_page_error", True):
                        raise
                    _LOG.warning("Skipped page %d: %s", page.id + 1, exc)

            _report(progress, 82, "Writing the Word file")
            target.parent.mkdir(parents=True, exist_ok=True)
            converter.make_docx(str(target), **settings)
        except PdfStudioError:
            raise
        except Exception as exc:
            if not _pdf_has_text_safely(source, password):
                raise PdfStudioError(
                    f"'{source.name}' has no selectable text, so it may be a scanned document. "
                    "Run OCR first, then convert it to Word."
                ) from exc
            raise PdfStudioError(f"'{source.name}' could not be converted to Word. {exc}") from exc
        finally:
            if converter is not None:
                with suppress(Exception):
                    converter.close()
    _report(progress, 100, "Done")
    return target


def _pdf_has_text_safely(path: Path, password: str | None) -> bool:
    try:
        return pdf_has_text(path, password)
    except Exception:
        return True  # Unknown: do not blame scanning for an unrelated failure.


# ---------------------------------------------------------------------------
# PDF to Excel
# ---------------------------------------------------------------------------


def pdf_to_xlsx(
    path: Path,
    output: Path,
    password: str | None = None,
    progress: ProgressCallback | None = None,
) -> Path:
    """Write one worksheet "Page N" per PDF page.

    Detected tables are written from the top with a blank row between them, and the
    first row of each table is bold. A page without tables gets its text lines in
    column A instead. Column widths are set from the content.
    """
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    source = Path(path)
    target = Path(output)
    _open_pdf(source, password).close()

    workbook = Workbook()
    default_sheet = workbook.active
    try:
        with pdfplumber.open(str(source), password=password or "") as pdf:
            pages = list(pdf.pages)
            total = len(pages)
            for number, page in enumerate(pages, start=1):
                _report(progress, _share(number - 1, total, 0, 90), f"Reading page {number} of {total}")
                sheet = workbook.create_sheet(f"Page {number}")
                _fill_sheet(sheet, page, Font, get_column_letter)
        if len(workbook.worksheets) > 1:
            workbook.remove(default_sheet)
        else:
            default_sheet.title = "Page 1"
        _report(progress, 94, "Saving the workbook")
        target.parent.mkdir(parents=True, exist_ok=True)
        with _remove_partial_output(target):
            workbook.save(str(target))
    except PdfStudioError:
        raise
    except Exception as exc:
        raise PdfStudioError(f"'{source.name}' could not be read as tables or text. {exc}") from exc
    _report(progress, 100, "Done")
    return target


def _fill_sheet(sheet: Any, page: Any, font_cls: Any, column_letter: Callable[[int], str]) -> None:
    widths: dict[int, int] = {}
    row = 1
    tables = [table for table in (page.extract_tables() or []) if _table_has_content(table)]
    if tables:
        for table_index, table in enumerate(tables):
            if table_index:
                row += 1  # blank row between tables
            columns = max(len(cells or []) for cells in table)
            for row_index, cells in enumerate(table):
                for col_index, value in enumerate(cells or [], start=1):
                    text = _clean_text(value)
                    if not text:
                        continue
                    _set_cell_text(sheet.cell(row=row, column=col_index), text)
                    widths[col_index] = max(widths.get(col_index, 0), _longest_line(text))
                if row_index == 0:
                    for col_index in range(1, columns + 1):
                        sheet.cell(row=row, column=col_index).font = font_cls(bold=True)
                row += 1
    else:
        for line in (page.extract_text() or "").splitlines():
            text = _clean_text(line)
            if text:
                _set_cell_text(sheet.cell(row=row, column=1), text)
                widths[1] = max(widths.get(1, 0), _longest_line(text))
            row += 1

    for col_index, width in widths.items():
        sheet.column_dimensions[column_letter(col_index)].width = min(60, max(8, width + 2))


def _table_has_content(table: Sequence[Sequence[Any]]) -> bool:
    return any(_clean_text(cell) for row in table for cell in (row or []))


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return _ILLEGAL_TEXT_CHARS.sub("", str(value)).strip()


def _longest_line(text: str) -> int:
    return max((len(line) for line in text.splitlines()), default=len(text))


def _set_cell_text(cell: Any, text: str) -> None:
    cell.value = text
    if text.startswith("="):
        cell.data_type = "s"  # keep PDF text such as "=5" as text, not a formula


# ---------------------------------------------------------------------------
# PDF to PowerPoint
# ---------------------------------------------------------------------------


def pdf_to_pptx(
    path: Path,
    output: Path,
    dpi: int = 150,
    pages: Sequence[int] | None = None,
    password: str | None = None,
    progress: ProgressCallback | None = None,
) -> Path:
    """One slide per page, sized to the page so the aspect ratio is kept.

    Each slide holds the rendered page as a picture on the Blank layout (no
    placeholders). The slides are image-based, not editable text.
    """
    from pptx import Presentation
    from pptx.util import Emu

    dpi = _check_dpi(dpi)
    source = Path(path)
    target = Path(output)
    document = _open_pdf(source, password)
    try:
        indexes = _resolve_pages(pages, document.page_count)
        first = document[indexes[0]].rect
        slide_w = int(round(first.width * _EMU_PER_PT))
        slide_h = int(round(first.height * _EMU_PER_PT))

        presentation = Presentation()
        presentation.slide_width = Emu(slide_w)
        presentation.slide_height = Emu(slide_h)
        blank_layout = presentation.slide_layouts[6]

        total = len(indexes)
        for step, index in enumerate(indexes):
            _report(progress, _share(step, total, 0, 95), f"Rendering page {index + 1} of {total}")
            page = document[index]
            picture = BytesIO(_render_jpeg(page, dpi))
            left, top, width, height = _picture_box(page.rect.width, page.rect.height, slide_w, slide_h)
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(picture, Emu(left), Emu(top), Emu(width), Emu(height))

        target.parent.mkdir(parents=True, exist_ok=True)
        with _remove_partial_output(target):
            presentation.save(str(target))
    finally:
        document.close()
    _report(progress, 100, "Done")
    return target


def _picture_box(page_w: float, page_h: float, slide_w: int, slide_h: int) -> tuple[int, int, int, int]:
    """Full-slide box for pages that match the deck size, otherwise a centred fit."""
    if abs(page_w * _EMU_PER_PT - slide_w) < 1 and abs(page_h * _EMU_PER_PT - slide_h) < 1:
        return 0, 0, slide_w, slide_h
    scale = min(slide_w / (page_w * _EMU_PER_PT), slide_h / (page_h * _EMU_PER_PT))
    width = int(round(page_w * _EMU_PER_PT * scale))
    height = int(round(page_h * _EMU_PER_PT * scale))
    return (slide_w - width) // 2, (slide_h - height) // 2, width, height


# ---------------------------------------------------------------------------
# HTML to PDF
# ---------------------------------------------------------------------------


def html_to_pdf(
    html: str,
    output: Path,
    page_size: str = "a4",
    margin: float = 36,
    progress: ProgressCallback | None = None,
) -> Path:
    """Lay out HTML with PyMuPDF's Story engine and write a PDF.

    Arabic text uses an Arabic-capable system font (Arial, Tahoma, DejaVu Sans or a
    Noto font) when one is found, so the text layer keeps the Arabic characters.
    If none is found the text is still converted.
    """
    return _html_to_pdf(html, Path(output), page_size, margin, progress, base_dir=None)


def html_file_to_pdf(
    path: Path,
    output: Path,
    page_size: str = "a4",
    margin: float = 36,
    progress: ProgressCallback | None = None,
) -> Path:
    """Read an .html file (UTF-8, falling back to legacy Arabic/Western code pages) and convert it.

    Images referenced by relative paths in the file are read from its folder.
    """
    source = Path(path)
    if not source.is_file():
        raise PdfStudioError(f"'{source.name}' was not found.")
    text = _decode_html(source.read_bytes())
    return _html_to_pdf(text, Path(output), page_size, margin, progress, base_dir=source.parent)


def _html_to_pdf(
    html: str,
    output: Path,
    page_size: str,
    margin: float,
    progress: ProgressCallback | None,
    base_dir: Path | None,
) -> Path:
    if not isinstance(html, str) or not html.strip():
        raise PdfStudioError("The HTML is empty. Add some content and try again.")
    size_key = str(page_size).lower()
    if size_key not in ("a4", "letter"):
        raise PdfStudioError("The page size must be A4 or Letter.")
    margin = float(margin)
    mediabox = pymupdf.paper_rect(size_key)
    where = mediabox + (margin, margin, -margin, -margin)
    if margin < 0 or where.width <= 0 or where.height <= 0:
        raise PdfStudioError("The margin is too large for the page size.")

    _report(progress, 0, "Preparing the HTML")
    css, archive = _html_styles(html, base_dir)
    story = pymupdf.Story(html=html, user_css=css, archive=archive)

    output.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(prefix=".pdfstudio-", suffix=".pdf", dir=str(output.parent))
    os.close(handle)
    temp = Path(temp_name)
    try:
        writer = pymupdf.DocumentWriter(str(temp))
        try:
            more = True
            page_count = 0
            while more:
                if page_count >= _MAX_HTML_PAGES:
                    raise PdfStudioError("The HTML is too long to convert.")
                device = writer.begin_page(mediabox)
                more, _ = story.place(where)
                story.draw(device)
                writer.end_page()
                page_count += 1
                _report(progress, min(85, 5 + 4 * page_count), f"Laying out page {page_count}")
        finally:
            writer.close()
        _report(progress, 90, "Checking the text")
        _verify_html_text(temp, html)
        os.replace(temp, output)
    except BaseException:
        _discard([temp])
        raise
    _report(progress, 100, "Done")
    return output


def _html_styles(html: str, base_dir: Path | None) -> tuple[str | None, pymupdf.Archive | None]:
    archive = pymupdf.Archive(str(base_dir)) if base_dir is not None else None
    if not _ARABIC_CHARS.search(html):
        return None, archive
    font = _find_arabic_font()
    if font is None:
        return None, archive
    if archive is None:
        archive = pymupdf.Archive()
    # Register the font in memory so the document CSS can name it by a relative URL.
    archive.add(font.read_bytes(), _BASE_FONT_FILE)
    css = (
        f"@font-face {{ font-family: 'PDF Studio Base'; src: url('{_BASE_FONT_FILE}'); }} "
        "html, body { font-family: 'PDF Studio Base', serif; }"
    )
    return css, archive


def _arabic_font_candidates() -> list[Path]:
    windows_dir = Path(os.environ.get("WINDIR") or os.environ.get("SystemRoot") or r"C:\Windows")
    candidates = [windows_dir / "Fonts" / name for name in ("arial.ttf", "tahoma.ttf", "segoeui.ttf", "times.ttf")]
    candidates += [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf"),
    ]
    return candidates


def _find_arabic_font() -> Path | None:
    for candidate in _arabic_font_candidates():
        try:
            if candidate.is_file() and pymupdf.Font(fontfile=str(candidate)).has_glyph(0x0627):
                return candidate
        except Exception:
            continue
    return None


def _decode_html(raw: bytes) -> str:
    """Decode HTML bytes: BOM, then a declared charset, then UTF-8, then Arabic or Western Windows code pages."""
    if raw.startswith(codecs.BOM_UTF8):
        return raw.decode("utf-8-sig", errors="replace")
    declared = _META_CHARSET.search(raw[:4096])
    if declared:
        try:
            return raw.decode(declared.group(1).decode("ascii"))
        except (LookupError, UnicodeDecodeError):
            pass
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        pass
    return raw.decode(_legacy_encoding(raw), errors="replace")


def _legacy_encoding(raw: bytes) -> str:
    """Choose cp1256 (Arabic) or cp1252 (Western) by which one yields more plausible letters."""
    arabic = sum(1 for ch in raw.decode("cp1256", errors="replace") if "\u0600" <= ch <= "\u06ff")
    western = sum(1 for ch in raw.decode("cp1252", errors="replace") if ch > "\x7f" and ch.isalpha())
    return "cp1256" if arabic > western else "cp1252"


class _PlainTextParser(HTMLParser):
    _SKIPPED = frozenset({"script", "style", "head", "title"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIPPED:
            self._skip_depth += 1
        self.parts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIPPED and self._skip_depth:
            self._skip_depth -= 1
        self.parts.append(" ")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def _verify_html_text(pdf_path: Path, html: str) -> None:
    """Check that the words of the HTML made it into the PDF text layer."""
    parser = _PlainTextParser()
    parser.feed(html)
    wanted = {word.lower() for word in _LATIN_WORD.findall("".join(parser.parts))}
    if not wanted:
        return  # Nothing verifiable (for example, Arabic-only text).
    document = pymupdf.open(str(pdf_path))
    try:
        actual = "".join(page.get_text() for page in document)
    finally:
        document.close()
    present = {word.lower() for word in _LATIN_WORD.findall(unicodedata.normalize("NFKC", actual))}
    if len(wanted & present) < 0.5 * len(wanted):
        raise PdfStudioError(
            "The text of the HTML could not be written to the PDF. Check the HTML and try again."
        )


# ---------------------------------------------------------------------------
# Office documents to PDF
# ---------------------------------------------------------------------------


def office_backend() -> str | None:
    """Return ``"office"`` (Microsoft Office on Windows), ``"libreoffice"`` or ``None``."""
    if _office_com_modules() is not None and any(
        _office_app_registered(app) for app in _OFFICE_PROG_IDS
    ):
        return "office"
    if find_libreoffice() is not None:
        return "libreoffice"
    return None


def office_to_pdf(path: Path, output: Path, progress: ProgressCallback | None = None) -> Path:
    """Convert .doc/.docx/.odt/.rtf, .xls/.xlsx/.ods and .ppt/.pptx/.odp to PDF.

    Uses Microsoft Office through COM on Windows when available, otherwise
    LibreOffice. Raises ``MissingDependencyError`` when neither is installed.
    """
    source = Path(path)
    target = Path(output)
    suffix = source.suffix.lower()
    if suffix not in _OFFICE_APPS:
        raise PdfStudioError(f"'{source.name}' is not a supported Office file.")
    if not source.is_file():
        raise PdfStudioError(f"'{source.name}' was not found.")
    app_key = _OFFICE_APPS[suffix]
    target.parent.mkdir(parents=True, exist_ok=True)

    with _remove_partial_output(target):
        if _office_com_modules() is not None and _office_app_registered(app_key):
            _com_convert(source, target, app_key, progress)
        else:
            soffice = find_libreoffice()
            if soffice is None:
                raise MissingDependencyError(_OFFICE_MISSING)
            _libreoffice_convert(soffice, source, target, progress)
        if not target.is_file():
            raise PdfStudioError(f"'{source.name}' could not be converted to PDF.")
    _report(progress, 100, "Done")
    return target


def _office_com_modules() -> tuple[Any, Any] | None:
    """Return ``(win32com.client, pythoncom)`` on Windows when pywin32 is installed."""
    if sys.platform != "win32":
        return None
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return None
    return win32com.client, pythoncom


def _office_app_registered(app_key: str) -> bool:
    if sys.platform != "win32":
        return False
    try:
        import winreg
    except ImportError:
        return False
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, _OFFICE_PROG_IDS[app_key]):
            return True
    except OSError:
        return False


def _com_convert(source: Path, target: Path, app_key: str, progress: ProgressCallback | None) -> None:
    modules = _office_com_modules()
    if modules is None:
        raise MissingDependencyError(_OFFICE_MISSING)
    com_client, pythoncom = modules
    label = _OFFICE_LABELS[app_key]
    src = str(source.resolve())
    dst = str(target.resolve())

    pythoncom.CoInitialize()  # Required on worker threads
    app: Any = None
    document: Any = None
    try:
        _report(progress, 15, f"Starting {label}")
        try:
            # DispatchEx starts a separate instance. Dispatch would attach to a Word or Excel
            # window the user already has open, and Quit() below would close it.
            app = com_client.DispatchEx(_OFFICE_PROG_IDS[app_key])
            with suppress(Exception):
                app.Visible = False
            with suppress(Exception):
                app.DisplayAlerts = {"word": 0, "excel": False, "powerpoint": 1}[app_key]

            _report(progress, 35, f"Opening the file in {label}")
            # pywin32 does not accept keyword arguments for COM methods, so the
            # optional parameters are passed by position (see the Office object model).
            if app_key == "word":
                document = app.Documents.Open(src, False, True, False)  # ConfirmConversions, ReadOnly, AddToRecentFiles
                _report(progress, 60, "Saving as PDF")
                document.SaveAs2(dst, 17)  # wdFormatPDF
            elif app_key == "excel":
                document = app.Workbooks.Open(src, 0, True)  # UpdateLinks, ReadOnly
                _report(progress, 60, "Saving as PDF")
                document.ExportAsFixedFormat(0, dst)  # xlTypePDF
            else:
                document = app.Presentations.Open(src, True, False, False)  # ReadOnly, Untitled, WithWindow
                _report(progress, 60, "Saving as PDF")
                document.SaveAs(dst, 32)  # ppSaveAsPDF
        except PdfStudioError:
            raise
        except Exception as exc:
            raise PdfStudioError(
                f"{label} could not convert '{source.name}'. The file may be damaged, password "
                f"protected or open in another program. ({exc})"
            ) from exc
    finally:
        _close_com(app, document, app_key)
        app = None
        document = None
        pythoncom.CoUninitialize()


def _close_com(app: Any, document: Any, app_key: str) -> None:
    if document is not None:
        with suppress(Exception):
            if app_key == "powerpoint":
                document.Close()
            else:
                document.Close(False)
    if app is not None:
        with suppress(Exception):
            app.Quit()


def find_libreoffice() -> str | None:
    """Return the path of LibreOffice's ``soffice`` program, or ``None``."""
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for folder in _program_files_dirs():
        candidate = folder / "LibreOffice" / "program" / "soffice.exe"
        if candidate.is_file():
            return str(candidate)
    return None


def _program_files_dirs() -> list[Path]:
    return [
        Path(os.environ.get("ProgramFiles", r"C:\Program Files")),
        Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")),
    ]


def _libreoffice_convert(soffice: str, source: Path, target: Path, progress: ProgressCallback | None) -> None:
    work = Path(tempfile.mkdtemp(prefix="pdfstudio-lo-"))
    try:
        outdir = work / "out"
        outdir.mkdir()
        profile = work / "profile"  # A private profile lets parallel conversions run side by side
        command = [
            soffice,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(source.resolve()),
        ]
        _report(progress, 20, "Converting with LibreOffice")
        try:
            process = subprocess.Popen(
                command,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                **_no_window_kwargs(),
            )
        except OSError as exc:
            raise PdfStudioError(f"LibreOffice could not be started. {exc}") from exc

        started = time.monotonic()
        try:
            while True:
                try:
                    stdout, stderr = process.communicate(timeout=1)
                    break
                except subprocess.TimeoutExpired:
                    elapsed = time.monotonic() - started
                    if elapsed > _LIBREOFFICE_TIMEOUT:
                        raise PdfStudioError(
                            f"LibreOffice took longer than {_LIBREOFFICE_TIMEOUT // 60} minutes "
                            f"to convert '{source.name}' and was stopped."
                        )
                    _report(progress, min(85, 20 + elapsed / 2), "Converting with LibreOffice")
        finally:
            if process.poll() is None:
                process.kill()
                with suppress(Exception):
                    process.communicate(timeout=10)

        produced = outdir / f"{source.stem}.pdf"
        if not produced.is_file():
            detail = _last_line((stderr or stdout or b"").decode("utf-8", errors="replace"))
            raise PdfStudioError(
                f"LibreOffice could not convert '{source.name}'." + (f" {detail}" if detail else "")
            )
        _report(progress, 90, "Saving the PDF")
        shutil.move(str(produced), str(target))
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _last_line(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return lines[-1][:200] if lines else ""


# ---------------------------------------------------------------------------
# OCR
# ---------------------------------------------------------------------------


def find_tesseract() -> str | None:
    """Return the path of the Tesseract executable, or ``None``."""
    for name in ("tesseract", "tesseract.exe"):
        found = shutil.which(name)
        if found:
            return found
    for folder in _program_files_dirs():
        candidate = folder / "Tesseract-OCR" / "tesseract.exe"
        if candidate.is_file():
            return str(candidate)
    return None


def _prepare_tessdata(tesseract: str) -> None:
    """Point TESSDATA_PREFIX at a tessdata folder next to the executable when none is set."""
    if os.environ.get("TESSDATA_PREFIX"):
        return
    tessdata = Path(tesseract).parent / "tessdata"
    if tessdata.is_dir():
        os.environ["TESSDATA_PREFIX"] = str(tessdata)


def ocr_languages(tesseract_path: str | None = None) -> list[str]:
    """Return the installed Tesseract language codes, or ``[]`` if they cannot be listed."""
    executable = tesseract_path or find_tesseract()
    if not executable:
        return []
    _prepare_tessdata(executable)
    try:
        completed = subprocess.run(
            [executable, "--list-langs"],
            capture_output=True,
            timeout=30,
            check=False,
            **_no_window_kwargs(),
        )
    except (OSError, subprocess.SubprocessError):
        return []
    if completed.returncode != 0:
        return []
    text = (completed.stdout or b"").decode("utf-8", errors="replace")
    text += "\n" + (completed.stderr or b"").decode("utf-8", errors="replace")
    codes: list[str] = []
    for line in text.splitlines():
        code = line.strip()
        if not code or code.lower().startswith("list of available") or " " in code:
            continue
        if code not in codes:
            codes.append(code)
    return codes


def ocr_pdf(
    path: Path,
    output: Path,
    language: str = "eng",
    dpi: int = 300,
    password: str | None = None,
    progress: ProgressCallback | None = None,
) -> Path:
    """Make a searchable PDF: each page is rendered, read by Tesseract and merged.

    Uses the ``pytesseract`` package when it is installed, otherwise runs the
    ``tesseract`` executable directly. Raises ``MissingDependencyError`` when Tesseract
    cannot be found.
    """
    source = Path(path)
    target = Path(output)
    tesseract = find_tesseract()
    if tesseract is None:
        raise MissingDependencyError(_TESSERACT_MISSING)
    dpi = _check_dpi(dpi)
    language = (language or "").strip() or "eng"
    if not _LANGUAGE_CODE.fullmatch(language):
        raise PdfStudioError("The OCR language code is not valid. Use a code such as eng or eng+fra.")
    _prepare_tessdata(tesseract)

    document = _open_pdf(source, password)
    result = pymupdf.open()
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        pytesseract = _import_pytesseract()
        if pytesseract is not None:
            pytesseract.pytesseract.tesseract_cmd = tesseract
        total = document.page_count
        with tempfile.TemporaryDirectory(prefix="pdfstudio-ocr-") as work_name:
            work = Path(work_name)
            for index in range(total):
                _report(progress, _share(index, total, 0, 90), f"Reading page {index + 1} of {total}")
                image_path = work / f"page-{index + 1:05d}.png"
                pixmap = _render_pixmap(document[index], dpi)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                image.save(image_path, format="PNG", dpi=(dpi, dpi))  # DPI tells Tesseract the page size
                page_pdf = _ocr_page_pdf(image_path, language, index + 1, tesseract, pytesseract, work)
                page_document = pymupdf.open(stream=page_pdf, filetype="pdf")
                try:
                    result.insert_pdf(page_document)
                finally:
                    page_document.close()
            _report(progress, 92, "Saving the searchable PDF")
            with _remove_partial_output(target):
                result.save(str(target), garbage=3, deflate=True)
    finally:
        result.close()
        document.close()
    _report(progress, 100, "Done")
    return target


def _import_pytesseract() -> Any | None:
    try:
        import pytesseract
    except ImportError:
        return None
    return pytesseract


def _ocr_page_pdf(
    image_path: Path,
    language: str,
    page_number: int,
    tesseract: str,
    pytesseract: Any | None,
    work: Path,
) -> bytes:
    if pytesseract is not None:
        try:
            return bytes(pytesseract.image_to_pdf_or_hocr(str(image_path), extension="pdf", lang=language))
        except Exception as exc:
            # The raw message is in args; str(exc) would repr-escape the quotes inside it.
            message = " ".join(str(arg) for arg in exc.args) or str(exc)
            raise _tesseract_error(page_number, message) from exc

    out_base = work / image_path.stem
    command = [tesseract, str(image_path), str(out_base), "-l", language, "pdf"]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            timeout=_TESSERACT_PAGE_TIMEOUT,
            check=False,
            **_no_window_kwargs(),
        )
    except subprocess.TimeoutExpired as exc:
        raise PdfStudioError(f"Tesseract took too long to read page {page_number}.") from exc
    except OSError as exc:
        raise PdfStudioError(f"Tesseract could not be started. {exc}") from exc
    pdf_file = Path(f"{out_base}.pdf")
    if completed.returncode != 0 or not pdf_file.is_file():
        raise _tesseract_error(page_number, completed.stderr or completed.stdout)
    return pdf_file.read_bytes()


def _tesseract_error(page_number: int, output: bytes | str | None) -> PdfStudioError:
    text = output.decode("utf-8", errors="replace") if isinstance(output, bytes) else (output or "")
    missing = re.search(r"Failed loading language '([^']+)'", text)
    if missing:
        return PdfStudioError(
            f"Tesseract could not load the OCR language '{missing.group(1)}'. "
            "Check that its data file is installed in the tessdata folder."
        )
    detail = _last_line(text)
    return PdfStudioError(f"Tesseract could not read page {page_number}." + (f" {detail}" if detail else ""))


def pdf_has_text(path: Path, password: str | None = None) -> bool:
    """True when at least one page has selectable text (scanned PDFs usually do not)."""
    document = _open_pdf(Path(path), password)
    try:
        return any(page.get_text("text").strip() for page in document)
    finally:
        document.close()
