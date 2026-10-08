"""Tests for pdfstudio.core.convert.

Sample PDFs, images and workbooks are created inside each test. LibreOffice and
Tesseract tests are skipped when the program is not installed.
"""

from __future__ import annotations

import unicodedata
import zipfile
from pathlib import Path

import pymupdf
import pytest
from PIL import Image

from pdfstudio.core import convert
from pdfstudio.core.errors import (
    MissingDependencyError,
    OperationCancelled,
    PasswordRequiredError,
    PdfStudioError,
)


def _make_text_pdf(path: Path, pages=("Alpha page one", "Beta page two", "Gamma page three"),
                   size=(612, 792)) -> Path:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page(width=size[0], height=size[1])
        page.insert_text((72, 100), text, fontsize=18)
    document.save(str(path))
    document.close()
    return path


def _make_table_pdf(path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    table = Table([["Item", "Quantity", "Price"], ["Apples", "12", "3.50"], ["Pears", "7", "4.25"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    style = getSampleStyleSheet()["Normal"]
    story = [table, PageBreak(), Paragraph("Plain text on page two", style)]
    SimpleDocTemplate(str(path), pagesize=letter).build(story)
    return path


def _page_text(path: Path) -> str:
    document = pymupdf.open(str(path))
    try:
        return "".join(page.get_text() for page in document)
    finally:
        document.close()


def _has_tesseract_english() -> bool:
    return convert.find_tesseract() is not None and "eng" in convert.ocr_languages()


# ---------------------------------------------------------------------------
# PDF to images
# ---------------------------------------------------------------------------


def test_pdf_to_images_jpg_names_and_dpi(tmp_path):
    pdf = _make_text_pdf(tmp_path / "Sample.pdf")
    out_dir = tmp_path / "images"

    files = convert.pdf_to_images(pdf, out_dir, "Sample", fmt="jpg", dpi=72)

    assert [f.name for f in files] == [
        "Sample - page 001.jpg",
        "Sample - page 002.jpg",
        "Sample - page 003.jpg",
    ]
    with Image.open(files[0]) as image:
        assert image.format == "JPEG"
        assert image.size == (612, 792)


def test_pdf_to_images_png_subset_keeps_page_numbers_and_dpi(tmp_path):
    pdf = _make_text_pdf(tmp_path / "Sample.pdf")

    files = convert.pdf_to_images(pdf, tmp_path / "out", "Sample", fmt="png", dpi=150, pages=[0, 2])

    assert [f.name for f in files] == ["Sample - page 001.png", "Sample - page 003.png"]
    with Image.open(files[1]) as image:
        assert image.format == "PNG"
        assert image.size == (1275, 1650)


def test_pdf_to_images_avoids_overwriting_and_validates_dpi(tmp_path):
    pdf = _make_text_pdf(tmp_path / "Sample.pdf", pages=("only",))
    first = convert.pdf_to_images(pdf, tmp_path, "Sample", dpi=72)
    second = convert.pdf_to_images(pdf, tmp_path, "Sample", dpi=72)

    assert first[0].name == "Sample - page 001.jpg"
    assert second[0].name == "Sample - page 001 (2).jpg"
    with pytest.raises(PdfStudioError):
        convert.pdf_to_images(pdf, tmp_path, "Sample", dpi=20)
    with pytest.raises(PdfStudioError):
        convert.pdf_to_images(pdf, tmp_path, "Sample", dpi=700)


def test_pdf_to_images_cancel_removes_partial_files(tmp_path):
    pdf = _make_text_pdf(tmp_path / "Sample.pdf")
    calls = []

    def progress(percent, message):
        calls.append(percent)
        if len(calls) == 3:
            raise OperationCancelled("cancelled by test")

    with pytest.raises(OperationCancelled):
        convert.pdf_to_images(pdf, tmp_path / "out", "Sample", dpi=72, progress=progress)
    assert not list((tmp_path / "out").glob("*"))


def test_password_protected_pdf_needs_password(tmp_path):
    plain = _make_text_pdf(tmp_path / "plain.pdf", pages=("Secret text",))
    locked = tmp_path / "locked.pdf"
    document = pymupdf.open(str(plain))
    document.save(str(locked), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="open", owner_pw="owner")
    document.close()

    with pytest.raises(PasswordRequiredError):
        convert.pdf_to_images(locked, tmp_path / "out", "x", dpi=72)
    with pytest.raises(PasswordRequiredError):
        convert.pdf_to_images(locked, tmp_path / "out", "x", dpi=72, password="wrong")
    files = convert.pdf_to_images(locked, tmp_path / "out", "x", dpi=72, password="open")
    assert len(files) == 1


# ---------------------------------------------------------------------------
# Images to PDF
# ---------------------------------------------------------------------------


def test_images_to_pdf_fit_keeps_pixel_size_and_transparency(tmp_path):
    rgba = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
    rgba.paste((255, 0, 0, 255), (0, 0, 100, 100))  # left half red, right half transparent
    png = tmp_path / "alpha.png"
    rgba.save(png)
    output = tmp_path / "fit.pdf"

    convert.images_to_pdf([png], output, page_size="fit")

    document = pymupdf.open(str(output))
    try:
        assert document.page_count == 1
        page = document[0]
        assert page.rect.width == pytest.approx(150)  # 200 px at 96 dpi
        assert page.rect.height == pytest.approx(75)
        pixel = page.get_pixmap(dpi=72, alpha=False).pixel(120, 40)
        assert pixel == (255, 255, 255)  # transparent area composited onto white
    finally:
        document.close()


def test_images_to_pdf_a4_centres_and_handles_cmyk_jpeg(tmp_path):
    landscape = tmp_path / "wide.jpg"
    Image.new("RGB", (400, 200), (255, 0, 0)).save(landscape, quality=95)
    cmyk = tmp_path / "cmyk.jpg"
    Image.new("CMYK", (300, 300), (255, 0, 0, 0)).save(cmyk, "JPEG")
    output = tmp_path / "a4.pdf"

    convert.images_to_pdf([landscape, cmyk], output, page_size="a4", margin=36)

    document = pymupdf.open(str(output))
    try:
        assert document.page_count == 2
        first = document[0]
        assert (round(first.rect.width), round(first.rect.height)) == (842, 595)  # landscape A4
        bbox = first.get_image_info()[0]["bbox"]
        assert bbox[0] >= 36 - 0.5 and bbox[2] <= 842 - 36 + 0.5
        assert (bbox[0] + bbox[2]) / 2 == pytest.approx(842 / 2, abs=1)  # centred
        assert (bbox[1] + bbox[3]) / 2 == pytest.approx(595 / 2, abs=1)
        # The RGB JPEG is embedded unchanged.
        xref = first.get_images()[0][0]
        assert document.extract_image(xref)["image"] == landscape.read_bytes()
        # CMYK converts to cyan for pure cyan ink (sample the centre of the page).
        pixmap = document[1].get_pixmap(dpi=72, alpha=False)
        r, g, b = pixmap.pixel(pixmap.width // 2, pixmap.height // 2)
        assert r < 60 and g > 180 and b > 180
    finally:
        document.close()


def test_images_to_pdf_rejects_empty_and_unreadable_input(tmp_path):
    with pytest.raises(PdfStudioError):
        convert.images_to_pdf([], tmp_path / "x.pdf")
    broken = tmp_path / "broken.png"
    broken.write_bytes(b"not really a png")
    with pytest.raises(PdfStudioError, match="broken.png"):
        convert.images_to_pdf([broken], tmp_path / "x.pdf")
    assert not (tmp_path / "x.pdf").exists()


# ---------------------------------------------------------------------------
# PDF to Word, Excel, PowerPoint
# ---------------------------------------------------------------------------


def test_pdf_to_docx_contains_page_text(tmp_path):
    pdf = _make_text_pdf(tmp_path / "report.pdf", pages=("Quarterly report summary",))
    output = tmp_path / "report.docx"

    result = convert.pdf_to_docx(pdf, output)

    assert result == output and output.is_file()
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "Quarterly" in xml and "report" in xml and "summary" in xml


def test_pdf_to_xlsx_one_sheet_per_page_with_table(tmp_path):
    from openpyxl import load_workbook

    pdf = _make_table_pdf(tmp_path / "table.pdf")
    output = tmp_path / "table.xlsx"

    convert.pdf_to_xlsx(pdf, output)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Page 1", "Page 2"]
    sheet = workbook["Page 1"]
    values = {cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
    assert {"Item", "Quantity", "Apples", "12", "3.50", "Pears", "4.25"} <= values
    assert sheet["A1"].font.bold
    assert workbook["Page 2"]["A1"].value == "Plain text on page two"


def test_pdf_to_pptx_slides_keep_page_aspect(tmp_path):
    from pptx import Presentation

    document = pymupdf.open()
    document.new_page(width=400, height=300).insert_text((50, 50), "Slide one")
    document.new_page(width=300, height=300).insert_text((50, 50), "Slide two")
    pdf = tmp_path / "slides.pdf"
    document.save(str(pdf))
    document.close()
    output = tmp_path / "slides.pptx"

    convert.pdf_to_pptx(pdf, output, dpi=72)

    presentation = Presentation(str(output))
    assert len(presentation.slides) == 2
    assert presentation.slide_width == 400 * 12700
    assert presentation.slide_height == 300 * 12700
    assert presentation.slide_width / presentation.slide_height == pytest.approx(400 / 300)
    first, second = presentation.slides
    for slide in presentation.slides:
        assert len(slide.placeholders) == 0
    assert first.shapes[0].width == presentation.slide_width
    # The 300 x 300 page is centred on the 400 x 300 slide.
    assert second.shapes[0].left == (400 - 300) * 12700 // 2
    assert second.shapes[0].width == 300 * 12700


# ---------------------------------------------------------------------------
# HTML and Office
# ---------------------------------------------------------------------------


def test_html_to_pdf_writes_text(tmp_path):
    output = tmp_path / "page.pdf"

    convert.html_to_pdf("<h1>Quarterly Overview</h1><p>Sample paragraph text here.</p>", output)

    assert "Sample paragraph" in _page_text(output)


def test_html_to_pdf_rejects_empty_html(tmp_path):
    with pytest.raises(PdfStudioError):
        convert.html_to_pdf("   ", tmp_path / "empty.pdf")


def test_html_file_to_pdf_decodes_windows_1256_arabic(tmp_path):
    if convert._find_arabic_font() is None:
        pytest.skip("No Arabic-capable system font was found")
    arabic = "مرحبا بالعالم"
    html = (
        '<html><head><meta charset="windows-1256"></head>'
        f"<body><p>Hello sample</p><p dir=\"rtl\">{arabic}</p></body></html>"
    )
    source = tmp_path / "arabic.html"
    source.write_bytes(html.encode("cp1256"))
    output = tmp_path / "arabic.pdf"

    convert.html_file_to_pdf(source, output)

    text = unicodedata.normalize("NFKC", _page_text(output))
    assert "Hello" in text
    assert arabic in text


def test_html_to_pdf_without_arabic_font_still_converts(tmp_path, monkeypatch):
    monkeypatch.setattr(convert, "_arabic_font_candidates", lambda: [])
    output = tmp_path / "arabic-no-font.pdf"

    convert.html_to_pdf("<p>Hello sample</p><p>\u0645\u0631\u062d\u0628\u0627</p>", output)

    assert "Hello sample" in _page_text(output)


def test_office_to_pdf_xlsx_via_libreoffice(tmp_path):
    if convert.office_backend() is None:
        pytest.skip("Neither Microsoft Office nor LibreOffice is installed")
    from openpyxl import Workbook

    source = tmp_path / "budget.xlsx"
    workbook = Workbook()
    workbook.active.column_dimensions["A"].width = 30  # keep the label from being clipped
    workbook.active["A1"] = "Invoice total"
    workbook.active["B1"] = 42
    workbook.save(source)
    output = tmp_path / "budget.pdf"

    convert.office_to_pdf(source, output)

    assert output.is_file()
    assert "Invoice total" in _page_text(output)


def test_office_to_pdf_rejects_unknown_type(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("plain", encoding="utf-8")
    with pytest.raises(PdfStudioError):
        convert.office_to_pdf(source, tmp_path / "notes.pdf")


# ---------------------------------------------------------------------------
# OCR and text detection
# ---------------------------------------------------------------------------


def _make_image_only_pdf(path: Path) -> Path:
    text_pdf = _make_text_pdf(path.with_suffix(".src.pdf"), pages=("INVOICE 2024", "TOTAL DUE"))
    source = pymupdf.open(str(text_pdf))
    document = pymupdf.open()
    try:
        for page in source:
            pixmap = page.get_pixmap(dpi=200, alpha=False)
            new_page = document.new_page(width=page.rect.width, height=page.rect.height)
            new_page.insert_image(new_page.rect, stream=pixmap.tobytes("png"))
        document.save(str(path))
    finally:
        document.close()
        source.close()
    return path


def test_pdf_has_text_distinguishes_text_and_scans(tmp_path):
    text_pdf = _make_text_pdf(tmp_path / "text.pdf")
    scan_pdf = _make_image_only_pdf(tmp_path / "scan.pdf")

    assert convert.pdf_has_text(text_pdf) is True
    assert convert.pdf_has_text(scan_pdf) is False


def test_ocr_pdf_without_tesseract_raises(tmp_path, monkeypatch):
    monkeypatch.setattr(convert, "find_tesseract", lambda: None)
    source = _make_text_pdf(tmp_path / "in.pdf", pages=("text",))

    with pytest.raises(MissingDependencyError, match="Tesseract"):
        convert.ocr_pdf(source, tmp_path / "out.pdf")


@pytest.mark.parametrize("backend", ["pytesseract", "cli"])
def test_ocr_pdf_adds_text_layer(tmp_path, monkeypatch, backend):
    if not _has_tesseract_english():
        pytest.skip("Tesseract OCR with English data is not installed")
    if backend == "pytesseract":
        pytest.importorskip("pytesseract")
    else:
        monkeypatch.setattr(convert, "_import_pytesseract", lambda: None)
    scan = _make_image_only_pdf(tmp_path / "scan.pdf")
    output = tmp_path / "searchable.pdf"

    convert.ocr_pdf(scan, output, language="eng", dpi=300)

    text = _page_text(output).upper()
    assert "INVOICE" in text
    assert "TOTAL" in text
    assert convert.pdf_has_text(output)
